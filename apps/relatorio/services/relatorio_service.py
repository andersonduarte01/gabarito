"""
Relatório service — agrega dados de outros módulos e retorna contexto
ou bytes (PDF/xlsx) conforme o formato solicitado.

formato: 'html' → dict de contexto
         'pdf'  → bytes (WeasyPrint)
         'xlsx' → bytes (openpyxl)
"""
from decimal import Decimal


# ---------------------------------------------------------------------------
# Helpers internos
# ---------------------------------------------------------------------------

def _exportar_xlsx(dados, colunas, titulo):
    """Gera bytes de uma planilha .xlsx simples."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise RuntimeError('openpyxl não está instalado.')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo[:31]

    header_fill  = PatternFill('solid', fgColor='0D6EFD')
    header_font  = Font(bold=True, color='FFFFFF')
    header_align = Alignment(horizontal='center')

    ws.append(colunas)
    for cell in ws[1]:
        cell.fill  = header_fill
        cell.font  = header_font
        cell.alignment = header_align

    for row in dados:
        ws.append(row)

    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _exportar_pdf(template_name, context, base_url=None):
    """Renderiza um template como PDF via WeasyPrint."""
    try:
        from weasyprint import HTML
    except ImportError:
        raise RuntimeError('WeasyPrint não está instalado.')
    from django.template.loader import render_to_string
    html_str = render_to_string(template_name, context)
    return HTML(string=html_str, base_url=base_url).write_pdf()


# ---------------------------------------------------------------------------
# Desempenho por turma
# ---------------------------------------------------------------------------

def gerar_desempenho_turma(turma, periodo_letivo, formato='html'):
    from apps.boletim.models import ResultadoPeriodo, SituacaoPeriodo
    from apps.materia.models import MateriaSerieConfig
    from apps.aluno.models import MatriculaTurma, SituacaoMatricula

    matriculas = MatriculaTurma.objects.filter(
        turma=turma,
        situacao=SituacaoMatricula.MATRICULADO,
        ativo=True,
    ).values_list('aluno_id', flat=True)

    filtro = dict(
        aluno_id__in=matriculas,
        periodo_letivo=periodo_letivo,
    )
    resultados = ResultadoPeriodo.objects.filter(**filtro).select_related('materia')

    materias = {}
    for r in resultados:
        m = materias.setdefault(r.materia, {
            'materia': r.materia,
            'total': 0, 'aprovados': 0, 'recuperacao': 0,
            'reprovados': 0, 'em_curso': 0, 'notas': [],
        })
        m['total'] += 1
        if r.media_periodo is not None:
            m['notas'].append(r.media_periodo)
        if r.situacao_periodo == SituacaoPeriodo.APROVADO:
            m['aprovados'] += 1
        elif r.situacao_periodo == SituacaoPeriodo.RECUPERACAO:
            m['recuperacao'] += 1
        elif r.situacao_periodo == SituacaoPeriodo.REPROVADO:
            m['reprovados'] += 1
        else:
            m['em_curso'] += 1

    for m in materias.values():
        notas = m.pop('notas')
        m['media_turma'] = round(sum(notas) / len(notas), 2) if notas else None

    linhas = sorted(materias.values(), key=lambda x: x['materia'].nome)
    ctx = {
        'turma': turma,
        'periodo_letivo': periodo_letivo,
        'linhas': linhas,
        'total_alunos': len(matriculas),
    }

    if formato == 'html':
        return ctx
    if formato == 'xlsx':
        colunas = ['Matéria', 'Total', 'Aprovados', 'Recuperação', 'Reprovados', 'Em Curso', 'Média da Turma']
        dados = [
            [
                l['materia'].nome, l['total'], l['aprovados'],
                l['recuperacao'], l['reprovados'], l['em_curso'],
                str(l['media_turma']) if l['media_turma'] is not None else '',
            ]
            for l in linhas
        ]
        return _exportar_xlsx(dados, colunas, f'Desempenho {turma.nome}')
    if formato == 'pdf':
        return _exportar_pdf('relatorio/pdf/desempenho_turma.html', ctx)
    raise ValueError(f'Formato desconhecido: {formato}')


# ---------------------------------------------------------------------------
# Frequência por turma
# ---------------------------------------------------------------------------

def gerar_frequencia_turma(turma, periodo_letivo, formato='html'):
    from apps.frequencia.models import RegistroFrequencia, PresencaAluno
    from apps.aluno.models import MatriculaTurma, SituacaoMatricula

    matriculas = MatriculaTurma.objects.filter(
        turma=turma, situacao=SituacaoMatricula.MATRICULADO, ativo=True,
    ).select_related('aluno').order_by('aluno__nome_completo')

    registros_ids = RegistroFrequencia.objects.filter(
        turma=turma, periodo_letivo=periodo_letivo,
    ).values_list('id', flat=True)

    total_aulas = len(registros_ids)

    linhas = []
    for mat in matriculas:
        presencas = PresencaAluno.objects.filter(
            registro_id__in=registros_ids, aluno=mat.aluno,
        )
        total_p  = presencas.filter(presente=True).count()
        total_f  = presencas.filter(presente=False, justificado=False).count()
        total_fj = presencas.filter(presente=False, justificado=True).count()
        pct = round(total_p / total_aulas * 100, 1) if total_aulas else Decimal('100.0')
        linhas.append({
            'aluno': mat.aluno,
            'presencas': total_p,
            'faltas': total_f,
            'faltas_justificadas': total_fj,
            'percentual': pct,
        })

    ctx = {
        'turma': turma,
        'periodo_letivo': periodo_letivo,
        'linhas': linhas,
        'total_aulas': total_aulas,
    }
    if formato == 'html':
        return ctx
    if formato == 'xlsx':
        colunas = ['Aluno', 'Presenças', 'Faltas', 'Faltas Justif.', '% Frequência']
        dados = [[l['aluno'].nome_completo, l['presencas'], l['faltas'], l['faltas_justificadas'], str(l['percentual'])] for l in linhas]
        return _exportar_xlsx(dados, colunas, f'Frequência {turma.nome}')
    if formato == 'pdf':
        return _exportar_pdf('relatorio/pdf/frequencia_turma.html', ctx)
    raise ValueError(f'Formato desconhecido: {formato}')


# ---------------------------------------------------------------------------
# Frequência por aluno
# ---------------------------------------------------------------------------

def gerar_frequencia_aluno(aluno, ano_letivo, formato='html'):
    from apps.frequencia.models import RegistroFrequencia, PresencaAluno
    from apps.aluno.models import MatriculaTurma, SituacaoMatricula

    matricula = MatriculaTurma.objects.filter(
        aluno=aluno, ano_letivo=ano_letivo, ativo=True,
    ).select_related('turma').first()

    if not matricula:
        return {'aluno': aluno, 'erro': 'Sem matrícula ativa neste ano letivo.'}

    turma = matricula.turma
    registros = RegistroFrequencia.objects.filter(
        turma=turma, ano_letivo=ano_letivo,
    ).select_related('materia', 'periodo_letivo').order_by('data')

    por_materia = {}
    for reg in registros:
        mat_nome = reg.materia.nome if reg.materia else 'Geral'
        bloco = por_materia.setdefault(mat_nome, {'total': 0, 'presencas': 0, 'faltas': 0})
        bloco['total'] += 1
        try:
            p = PresencaAluno.objects.get(registro=reg, aluno=aluno)
            if p.presente:
                bloco['presencas'] += 1
            else:
                bloco['faltas'] += 1
        except PresencaAluno.DoesNotExist:
            pass

    for bloco in por_materia.values():
        bloco['percentual'] = round(bloco['presencas'] / bloco['total'] * 100, 1) if bloco['total'] else 100.0

    ctx = {'aluno': aluno, 'ano_letivo': ano_letivo, 'turma': turma, 'por_materia': por_materia}
    if formato == 'html':
        return ctx
    if formato == 'xlsx':
        colunas = ['Matéria', 'Total Aulas', 'Presenças', 'Faltas', '% Frequência']
        dados = [[mat, b['total'], b['presencas'], b['faltas'], str(b['percentual'])] for mat, b in por_materia.items()]
        return _exportar_xlsx(dados, colunas, f'Frequência {aluno.nome_completo}')
    if formato == 'pdf':
        return _exportar_pdf('relatorio/pdf/frequencia_aluno.html', ctx)
    raise ValueError(f'Formato desconhecido: {formato}')


# ---------------------------------------------------------------------------
# Alunos em risco
# ---------------------------------------------------------------------------

def gerar_alunos_em_risco(escola, periodo_letivo, formato='html'):
    from apps.boletim.models import ResultadoPeriodo, SituacaoPeriodo
    from apps.configuracao.models import ConfiguracaoFrequencia

    cfg = ConfiguracaoFrequencia.objects.filter(escola=escola).first()
    pct_minimo = cfg.percentual_minimo_frequencia if cfg else Decimal('75.0')

    resultados = ResultadoPeriodo.objects.filter(
        aluno__escola=escola,
        periodo_letivo=periodo_letivo,
        situacao_periodo__in=[SituacaoPeriodo.RECUPERACAO, SituacaoPeriodo.REPROVADO],
    ).select_related('aluno', 'materia').order_by('aluno__nome_completo', 'materia__nome')

    freq_risco = ResultadoPeriodo.objects.filter(
        aluno__escola=escola,
        periodo_letivo=periodo_letivo,
        frequencia_percentual__lt=pct_minimo,
    ).values_list('aluno_id', flat=True)

    alunos_freq_risco = set(freq_risco)

    por_aluno = {}
    for r in resultados:
        bloco = por_aluno.setdefault(r.aluno, {
            'aluno': r.aluno, 'materias_nota': [], 'risco_freq': r.aluno_id in alunos_freq_risco,
        })
        bloco['materias_nota'].append({'materia': r.materia, 'situacao': r.situacao_periodo, 'media': r.media_periodo})

    linhas = sorted(por_aluno.values(), key=lambda x: x['aluno'].nome_completo)

    ctx = {'periodo_letivo': periodo_letivo, 'linhas': linhas, 'total': len(linhas)}
    if formato == 'html':
        return ctx
    if formato == 'xlsx':
        colunas = ['Aluno', 'Matéria', 'Situação', 'Média', 'Risco de Freq.']
        dados = []
        for bloco in linhas:
            for m in bloco['materias_nota']:
                dados.append([
                    bloco['aluno'].nome_completo,
                    m['materia'].nome,
                    m['situacao'],
                    str(m['media']) if m['media'] else '',
                    'Sim' if bloco['risco_freq'] else 'Não',
                ])
        return _exportar_xlsx(dados, colunas, 'Alunos em Risco')
    if formato == 'pdf':
        return _exportar_pdf('relatorio/pdf/alunos_em_risco.html', ctx)
    raise ValueError(f'Formato desconhecido: {formato}')


# ---------------------------------------------------------------------------
# Boletim em lote (PDF apenas — WeasyPrint)
# ---------------------------------------------------------------------------

def gerar_boletim_lote(turma, ano_letivo, formato='pdf'):
    from apps.aluno.models import MatriculaTurma, SituacaoMatricula
    from apps.boletim.models import ResultadoPeriodo, ResultadoAnual
    from apps.ano_letivo.models import PeriodoLetivo

    matriculas = MatriculaTurma.objects.filter(
        turma=turma, ano_letivo=ano_letivo,
        situacao=SituacaoMatricula.MATRICULADO, ativo=True,
    ).select_related('aluno').order_by('aluno__nome_completo')

    periodos = PeriodoLetivo.objects.filter(ano_letivo=ano_letivo).order_by('numero')

    boletins = []
    for mat in matriculas:
        aluno = mat.aluno
        resultados_p = ResultadoPeriodo.objects.filter(
            aluno=aluno, ano_letivo=ano_letivo,
        ).select_related('materia', 'periodo_letivo').order_by('materia__nome', 'periodo_letivo__numero')

        materias_map = {}
        for r in resultados_p:
            m = materias_map.setdefault(r.materia.nome, {})
            m[r.periodo_letivo.numero] = r

        resultado_anual = ResultadoAnual.objects.filter(aluno=aluno, ano_letivo=ano_letivo).select_related('materia')
        anual_map = {r.materia.nome: r for r in resultado_anual}

        boletins.append({
            'aluno': aluno,
            'materias_map': materias_map,
            'anual_map': anual_map,
        })

    ctx = {
        'turma': turma,
        'ano_letivo': ano_letivo,
        'periodos': periodos,
        'boletins': boletins,
        'escola': turma.escola,
    }

    if formato == 'html':
        return ctx
    if formato == 'pdf':
        return _exportar_pdf('relatorio/pdf/boletim_lote.html', ctx)
    raise ValueError(f'Formato desconhecido: {formato}')


# ---------------------------------------------------------------------------
# Desempenho por professor (DIRETOR only)
# ---------------------------------------------------------------------------

def gerar_desempenho_professor(escola, periodo_letivo, formato='html'):
    from apps.turma.models import ProfessorMateriaTurma
    from apps.boletim.models import ResultadoPeriodo, SituacaoPeriodo
    from apps.aluno.models import MatriculaTurma, SituacaoMatricula
    from apps.configuracao.models import ConfiguracaoAcademica

    cfg = ConfiguracaoAcademica.objects.filter(escola=escola).first()
    nota_minima = cfg.nota_minima_aprovacao if cfg else Decimal('5.0')

    vinculos = ProfessorMateriaTurma.objects.filter(
        turma__escola=escola, ativo=True,
    ).select_related('professor__papel__vinculo__usuario', 'materia', 'turma')

    por_professor = {}
    for v in vinculos:
        prof = v.professor
        nome_prof = prof.papel.vinculo.usuario.get_full_name() or prof.papel.vinculo.usuario.email
        bloco = por_professor.setdefault(nome_prof, {'professor': prof, 'nome': nome_prof, 'turmas': []})

        matriculas_ids = MatriculaTurma.objects.filter(
            turma=v.turma, situacao=SituacaoMatricula.MATRICULADO, ativo=True,
        ).values_list('aluno_id', flat=True)

        resultados = ResultadoPeriodo.objects.filter(
            aluno_id__in=matriculas_ids,
            materia=v.materia,
            periodo_letivo=periodo_letivo,
        )
        total    = resultados.count()
        notas    = [r.media_periodo for r in resultados if r.media_periodo is not None]
        abaixo   = sum(1 for n in notas if n < nota_minima)
        media_t  = round(sum(notas) / len(notas), 2) if notas else None
        pct_abaixo = round(abaixo / total * 100, 1) if total else 0

        bloco['turmas'].append({
            'turma': v.turma,
            'materia': v.materia,
            'total': total,
            'media': media_t,
            'abaixo_minimo': abaixo,
            'pct_abaixo': pct_abaixo,
            'alerta': pct_abaixo >= 50,
        })

    linhas = sorted(por_professor.values(), key=lambda x: x['nome'])

    ctx = {'periodo_letivo': periodo_letivo, 'linhas': linhas, 'nota_minima': nota_minima}
    if formato == 'html':
        return ctx
    if formato == 'xlsx':
        colunas = ['Professor', 'Turma', 'Matéria', 'Total Alunos', 'Média', 'Abaixo do Mínimo', '% Abaixo', 'Alerta']
        dados = []
        for bloco in linhas:
            for t in bloco['turmas']:
                dados.append([
                    bloco['nome'], t['turma'].nome, t['materia'].nome,
                    t['total'], str(t['media']) if t['media'] else '',
                    t['abaixo_minimo'], str(t['pct_abaixo']),
                    'SIM' if t['alerta'] else '',
                ])
        return _exportar_xlsx(dados, colunas, 'Desempenho Professor')
    if formato == 'pdf':
        return _exportar_pdf('relatorio/pdf/desempenho_professor.html', ctx)
    raise ValueError(f'Formato desconhecido: {formato}')


# ---------------------------------------------------------------------------
# Inadimplência
# ---------------------------------------------------------------------------

def gerar_inadimplencia(escola, data_inicio=None, data_fim=None, turma=None, formato='html'):
    from apps.financeiro.models import CobrancaAluno, StatusCobranca

    qs = CobrancaAluno.objects.filter(
        aluno__escola=escola,
        status=StatusCobranca.VENCIDO,
    ).select_related('aluno', 'responsavel_financeiro', 'plano_financeiro').order_by('vencimento')

    if data_inicio:
        qs = qs.filter(vencimento__gte=data_inicio)
    if data_fim:
        qs = qs.filter(vencimento__lte=data_fim)
    if turma:
        from apps.aluno.models import MatriculaTurma
        alunos_ids = MatriculaTurma.objects.filter(turma=turma, ativo=True).values_list('aluno_id', flat=True)
        qs = qs.filter(aluno_id__in=alunos_ids)

    total_valor = sum(c.valor for c in qs)

    ctx = {
        'cobrancas': qs,
        'total_cobrado': total_valor,
        'total_registros': qs.count(),
        'data_inicio': data_inicio,
        'data_fim': data_fim,
        'turma': turma,
    }
    if formato == 'html':
        return ctx
    if formato == 'xlsx':
        colunas = ['Aluno', 'Responsável', 'Descrição', 'Valor', 'Vencimento']
        dados = [
            [
                c.aluno.nome_completo,
                c.responsavel_financeiro.nome if c.responsavel_financeiro else '',
                c.descricao,
                str(c.valor),
                c.vencimento.strftime('%d/%m/%Y'),
            ]
            for c in qs
        ]
        return _exportar_xlsx(dados, colunas, 'Inadimplência')
    if formato == 'pdf':
        return _exportar_pdf('relatorio/pdf/inadimplencia.html', ctx)
    raise ValueError(f'Formato desconhecido: {formato}')


# ---------------------------------------------------------------------------
# Extrato financeiro
# ---------------------------------------------------------------------------

def gerar_extrato_financeiro(escola, data_inicio, data_fim, formato='html'):
    from apps.financeiro.models import CobrancaAluno, StatusCobranca

    qs = CobrancaAluno.objects.filter(
        aluno__escola=escola,
        vencimento__gte=data_inicio,
        vencimento__lte=data_fim,
    ).select_related('aluno', 'plano_financeiro').order_by('vencimento', 'status')

    totais = {}
    for status in StatusCobranca:
        totais[status.value] = {'label': status.label, 'qtd': 0, 'valor': Decimal('0')}

    cobrancas = list(qs)
    for c in cobrancas:
        totais[c.status]['qtd'] += 1
        totais[c.status]['valor'] += c.valor

    ctx = {
        'cobrancas': cobrancas,
        'totais': totais,
        'data_inicio': data_inicio,
        'data_fim': data_fim,
        'total_geral': sum(t['valor'] for t in totais.values()),
    }
    if formato == 'html':
        return ctx
    if formato == 'xlsx':
        colunas = ['Aluno', 'Descrição', 'Valor', 'Vencimento', 'Status', 'Pago em']
        dados = [
            [
                c.aluno.nome_completo, c.descricao, str(c.valor),
                c.vencimento.strftime('%d/%m/%Y'),
                c.get_status_display(),
                c.pago_em.strftime('%d/%m/%Y') if c.pago_em else '',
            ]
            for c in cobrancas
        ]
        return _exportar_xlsx(dados, colunas, 'Extrato Financeiro')
    if formato == 'pdf':
        return _exportar_pdf('relatorio/pdf/extrato_financeiro.html', ctx)
    raise ValueError(f'Formato desconhecido: {formato}')


# ---------------------------------------------------------------------------
# Executar relatórios agendados (chamado pelo management command)
# ---------------------------------------------------------------------------

def executar_agendados(escola):
    from apps.relatorio.models import RelatorioAgendado, RelatorioGerado, TipoRelatorio
    from apps.notificacao.models import TipoNotificacao
    import apps.notificacao.services.notificacao_service as notif_svc
    from apps.core.models import VinculoEscola
    import io, os
    from django.core.files.base import ContentFile

    agendados = RelatorioAgendado.objects.filter(escola=escola, ativo=True)
    gerados = 0

    for ag in agendados:
        try:
            ctx = _executar_um(ag)
            if ctx is None:
                continue
            # Salva como arquivo xlsx por padrão para agendados
            nome_arquivo = f'{ag.get_tipo_display().lower().replace(" ", "_")}.xlsx'
            bytes_arquivo = None

            if ag.tipo == TipoRelatorio.BOLETIM_LOTE:
                bytes_arquivo = None  # boletim lote só faz sentido sob demanda com turma
            else:
                # tenta gerar xlsx a partir do context — re-executa com formato xlsx
                try:
                    bytes_arquivo = _executar_um(ag, formato='xlsx')
                except Exception:
                    bytes_arquivo = None

            rel = RelatorioGerado.objects.create(
                agendado=ag,
                escola=escola,
                tipo=ag.tipo,
                parametros=ag.parametros,
            )
            if bytes_arquivo:
                rel.arquivo.save(nome_arquivo, ContentFile(bytes_arquivo), save=True)

            # Notifica os diretores da escola
            vinculos = VinculoEscola.objects.filter(
                escola=escola, ativo=True, papeis__tipo='DIRETOR', papeis__ativo=True,
            ).select_related('usuario').distinct()
            for v in vinculos:
                notif_svc.criar(
                    destinatario=v.usuario,
                    titulo=f'Relatório gerado: {ag.get_tipo_display()}',
                    mensagem='Seu relatório agendado foi gerado e está disponível.',
                    tipo=TipoNotificacao.SISTEMA,
                )
            gerados += 1
        except Exception:
            pass

    return gerados


def _executar_um(agendado, formato='html'):
    from apps.relatorio.models import TipoRelatorio
    from apps.turma.models import Turma
    from apps.ano_letivo.models import AnoLetivo, PeriodoLetivo

    p  = agendado.parametros
    tp = agendado.tipo

    turma   = Turma.objects.filter(id=p.get('turma_id')).first() if p.get('turma_id') else None
    periodo = PeriodoLetivo.objects.filter(id=p.get('periodo_id')).first() if p.get('periodo_id') else None
    ano     = AnoLetivo.objects.filter(id=p.get('ano_id')).first() if p.get('ano_id') else None

    if tp == TipoRelatorio.DESEMPENHO_TURMA and turma and periodo:
        return gerar_desempenho_turma(turma, periodo, formato)
    if tp == TipoRelatorio.FREQUENCIA_TURMA and turma and periodo:
        return gerar_frequencia_turma(turma, periodo, formato)
    if tp == TipoRelatorio.ALUNOS_EM_RISCO and periodo:
        return gerar_alunos_em_risco(agendado.escola, periodo, formato)
    if tp == TipoRelatorio.INADIMPLENCIA:
        return gerar_inadimplencia(agendado.escola, formato=formato)
    return None
